import os
import openpyxl
from django.core.management.base import BaseCommand, CommandError
from farmerapp.models import SoilType, IrrigationType, CropType

DEFAULT_EXCEL_PATH = os.path.join(
    os.path.dirname(__file__), "..", "..", "data", "translations.xlsx"
)


LANGUAGE_COLUMN_MAP = {
    "Hindi": "name_hi",
    "Marathi": "name_mr",
    "Gujarati": "name_gu",
    "Punjabi": "name_pa",
    "Bangla": "name_bn",
    "Kannada": "name_kn",
    "Telugu": "name_te",
    "Tamil": "name_ta",
}

SHEET_CONFIG = [
    {
        "sheet": "Crops",
        "model": CropType,
        "english_col": "crop_name (english)",
    },
    {
        "sheet": "Irrigation Type",
        "model": IrrigationType,
        "english_col": "irrigation_type_name (english)",
    },
    {
        "sheet": "Soil Type",
        "model": SoilType,
        "english_col": "soil_type_name (english)",
    },
]


class Command(BaseCommand):
    help = "Seed translations for CropType, IrrigationType, and SoilType from an Excel file"

    def add_arguments(self, parser):
        parser.add_argument(
            "--excel-path",
            type=str,
            default=DEFAULT_EXCEL_PATH,
            help="Path to the translations Excel file (default: farmerapp/data/translations.xlsx)",
        )

    def handle(self, *args, **options):
        excel_path = options["excel_path"]

        try:
            wb = openpyxl.load_workbook(excel_path)
        except FileNotFoundError:
            raise CommandError(f"File not found: {excel_path}")
        except Exception as e:
            raise CommandError(f"Failed to open Excel file: {e}")

        for config in SHEET_CONFIG:
            sheet_name = config["sheet"]
            model = config["model"]
            english_col = config["english_col"]
            model_name = model.__name__

            if sheet_name not in wb.sheetnames:
                self.stdout.write(self.style.WARNING(f"Sheet '{sheet_name}' not found, skipping."))
                continue

            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))

            if not rows:
                self.stdout.write(self.style.WARNING(f"Sheet '{sheet_name}' is empty, skipping."))
                continue

            headers = [str(h).strip() if h is not None else None for h in rows[0]]

            try:
                id_idx = headers.index("id")
                english_idx = headers.index(english_col)
            except ValueError as e:
                raise CommandError(f"Expected column not found in sheet '{sheet_name}': {e}")

            lang_indices = {}
            for col_header, field_name in LANGUAGE_COLUMN_MAP.items():
                if col_header in headers:
                    lang_indices[field_name] = headers.index(col_header)

            created_count = 0
            updated_count = 0
            skipped_count = 0

            for row in rows[1:]:
                irriwatch_id = row[id_idx]
                english_name = row[english_idx]

                if irriwatch_id is None or english_name is None:
                    skipped_count += 1
                    continue

                irriwatch_id = int(irriwatch_id)
                english_name = str(english_name).strip()

                defaults = {"name": english_name}
                for field_name, col_idx in lang_indices.items():
                    value = row[col_idx]
                    defaults[field_name] = str(value).strip() if value is not None else ""

                _, created = model.objects.update_or_create(
                    irriwatch_id=irriwatch_id,
                    defaults=defaults,
                )
                if created:
                    created_count += 1
                else:
                    updated_count += 1

            self.stdout.write(self.style.SUCCESS(
                f"{model_name}: {created_count} created, {updated_count} updated, {skipped_count} skipped"
            ))

        self.stdout.write(self.style.SUCCESS("Translation seeding complete."))
